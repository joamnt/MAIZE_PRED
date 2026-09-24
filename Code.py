# ==========================================================
# 1. IMPORTAÇÃO DAS BIBLIOTECAS
# ==========================================================
import pandas as pd
import numpy as np
import math
import warnings
from sklearn.model_selection import KFold, RandomizedSearchCV, GridSearchCV
from sklearn.preprocessing import StandardScaler, OneHotEncoder
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.metrics import r2_score, mean_squared_error, mean_absolute_error
from sklearn.base import clone
from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import LinearRegression
from sklearn.cross_decomposition import PLSRegression
from tensorflow.keras.models import Sequential
from tensorflow.keras.layers import Dense, Input
from scikeras.wrappers import KerasRegressor
import xgboost as xgb
from docx import Document

warnings.filterwarnings("ignore")


# ==========================================================
# 2. INICIALIZAÇÃO DO DOCUMENTO WORD E CONFIGURAÇÕES
# ==========================================================
document = Document()
document.add_heading('Relatório de Resultados - Modelagem Preditiva', level=1)

FILE_PATH = "/mnt/5a3c0750-b5d9-4421-aa5f-cd601ca4e1de/Arquivos i9/Milho e Soja/Milho.txt"
FACTOR_COLS = ["Safra", "Variedade", "N", "Bloco"]
TARGET_COLS = ["PROD"]
NON_PREDICTOR_COLS = FACTOR_COLS + ["AIE", "ALT", "NF"]


# ==========================================================
# 3. FUNÇÕES AUXILIARES
# ==========================================================
def create_mlp(
    meta,
    n_hidden1=20,
    n_hidden2=10,
    n_hidden3=0,
    activation="relu",
    optimizer="adam",
    learning_rate=0.001
):
    n_features_in = meta["n_features_in_"]

    from tensorflow.keras.optimizers import Adam, RMSprop

    opt = (
        Adam(learning_rate=learning_rate)
        if optimizer == "adam"
        else RMSprop(learning_rate=learning_rate)
    )

    model = Sequential([
        Input(shape=(n_features_in,)),
        Dense(n_hidden1, activation=activation)
    ])

    if n_hidden2 > 0:
        model.add(Dense(n_hidden2, activation=activation))

    if n_hidden3 > 0:
        model.add(Dense(n_hidden3, activation=activation))

    model.add(Dense(1))
    model.compile(loss="mse", optimizer=opt)

    return model


def calc_metrics(y_true, y_pred):
    y_true = np.array(y_true).ravel()
    y_pred = np.array(y_pred).ravel()

    if len(y_true) < 2 or np.std(y_pred) == 0:
        return {
            "r": 0,
            "R2": 0,
            "RMSE": 0,
            "MAE": 0
        }

    return {
        "r": np.corrcoef(y_true, y_pred)[0, 1],
        "R2": r2_score(y_true, y_pred),
        "RMSE": np.sqrt(mean_squared_error(y_true, y_pred)),
        "MAE": mean_absolute_error(y_true, y_pred)
    }


def adicionar_dataframe_ao_doc(doc, df):
    tabela = doc.add_table(rows=1, cols=len(df.columns))
    tabela.style = 'Table Grid'

    hdr_cells = tabela.rows[0].cells

    for i, col_name in enumerate(df.columns):
        hdr_cells[i].text = str(col_name)

    for _, row in df.iterrows():
        row_cells = tabela.add_row().cells

        for i, value in enumerate(row):
            row_cells[i].text = (
                f"{value:.4f}"
                if isinstance(value, (int, float))
                else str(value)
            )

    doc.add_paragraph()


# ==========================================================
# 4. EXECUÇÃO DA ANÁLISE
# ==========================================================
try:
    df = pd.read_csv(FILE_PATH, sep="\t")
except Exception as e:
    print(f"Erro: {e}")
    exit()


for col in FACTOR_COLS:
    df[col] = df[col].astype("category")


safras = list(df["Safra"].unique())

res_m1, res_m2, res_m3, res_m4 = [], [], [], []

todas_observacoes = []

best_params_mlp = {}
best_params_rf = {}
best_params_pls = {}


for y_col in TARGET_COLS:

    print(f"\n>>> Processando: {y_col}")

    X = pd.concat(
        [
            df[FACTOR_COLS],
            df.drop(
                columns=NON_PREDICTOR_COLS + TARGET_COLS,
                errors="ignore"
            )
        ],
        axis=1
    )

    y_target = df[y_col]


    # ======================================================
    # PREPROCESSAMENTO
    # ======================================================
    preprocessor = ColumnTransformer([
        (
            "cat",
            OneHotEncoder(
                handle_unknown="ignore",
                sparse_output=False
            ),
            X.select_dtypes(["category"]).columns
        ),
        (
            "num",
            StandardScaler(),
            X.select_dtypes(exclude=["category"]).columns
        )
    ])


    # ======================================================
    # ESCALONAMENTO DE Y PARA OTIMIZAÇÃO
    # ======================================================
    scaler_y_opt = StandardScaler()

    y_scaled_opt = scaler_y_opt.fit_transform(
        y_target.values.reshape(-1, 1)
    ).ravel()


    # ======================================================
    # CROSS-VALIDATION
    # ======================================================
    kf = KFold(
        n_splits=5,
        shuffle=True,
        random_state=42
    )


    # ======================================================
    # MLP SEARCH
    # ======================================================
    mlp_search = KerasRegressor(
        model=create_mlp,
        verbose=0,
        n_hidden1=32,
        n_hidden2=16,
        optimizer="adam",
        learning_rate=0.001
    )

    random_search_mlp = RandomizedSearchCV(
        Pipeline([
            ("pre", preprocessor),
            ("m", mlp_search)
        ]),
        {
            "m__n_hidden1": [32, 64],
            "m__learning_rate": [0.01, 0.001],
            "m__epochs": [50]
        },
        n_iter=5,
        cv=kf,
        n_jobs=-1
    ).fit(X, y_scaled_opt)

    best_params_mlp[y_col] = random_search_mlp.best_params_


    # ======================================================
    # RF SEARCH
    # ======================================================
    rf_search = RandomForestRegressor(
        random_state=42
    )

    random_search_rf = RandomizedSearchCV(
        Pipeline([
            ("pre", preprocessor),
            ("m", rf_search)
        ]),
        {
            "m__n_estimators": [100, 200],
            "m__max_depth": [10, None]
        },
        n_iter=5,
        cv=kf,
        n_jobs=-1
    ).fit(X, y_scaled_opt)

    best_params_rf[y_col] = random_search_rf.best_params_


    # ======================================================
    # PLS SEARCH
    # ======================================================
    pls_search = PLSRegression()

    grid_search_pls = GridSearchCV(
        Pipeline([
            ("pre", preprocessor),
            ("m", pls_search)
        ]),
        {
            "m__n_components": list(range(1, 21))
        },
        cv=kf,
        scoring="r2",
        n_jobs=-1
    ).fit(X, y_scaled_opt)

    best_params_pls[y_col] = grid_search_pls.best_params_


    # Mostrar o melhor número de componentes
    print(
        f"  Melhor n_components para PLS: "
        f"{grid_search_pls.best_params_['m__n_components']}"
    )


    # ======================================================
    # MODELOS
    # ======================================================
    modelos = {

        "XGBoost": Pipeline([
            ("pre", preprocessor),
            ("m", xgb.XGBRegressor(
                random_state=42
            ))
        ]),

        "MLP": random_search_mlp.best_estimator_,

        "Random Forest": random_search_rf.best_estimator_,

        "PLS": grid_search_pls.best_estimator_,

        "Linear Regression": Pipeline([
            ("pre", preprocessor),
            ("m", LinearRegression())
        ])
    }


    # ======================================================
    # ANÁLISES POR SAFRA
    # ======================================================
    for s in safras:

        print(f"  Analise Safra: {s}")

        df_s = df[df["Safra"] == s].copy()

        X_s = df_s.drop(columns=TARGET_COLS)
        y_s = df_s[y_col]


        # ==================================================
        # M1: Baseline
        # ==================================================
        scaler_y_m1 = StandardScaler()

        y_s_sc = scaler_y_m1.fit_transform(
            y_s.values.reshape(-1, 1)
        ).ravel()


        for nome, pipe in modelos.items():

            m = clone(pipe).fit(
                X_s,
                y_s_sc
            )

            pred = scaler_y_m1.inverse_transform(
                m.predict(X_s).reshape(-1, 1)
            ).ravel()


            res_m1.append({
                "Variável": y_col,
                "Safra": s,
                "Modelo": nome,
                **calc_metrics(y_s, pred)
            })


            for r, p in zip(y_s, pred):

                todas_observacoes.append({
                    "Metodo": "M1",
                    "Var": y_col,
                    "Safra": s,
                    "Mod": nome,
                    "Cen": "Base",
                    "Rep": 1,
                    "Real": r,
                    "Pred": p
                })


        # ==================================================
        # M2: Aleatório (1-4 blocos removidos)
        # ==================================================
        for n_rem in [1, 2, 3, 4]:

            metrics_rep = {
                m: []
                for m in modelos
            }


            for rep in range(50):

                t_idx = (
                    df_s
                    .groupby(
                        "Variedade",
                        group_keys=False
                    )
                    .sample(
                        n=n_rem,
                        random_state=rep * 100
                    )
                    .index
                )


                tr_idx = df_s.index.difference(
                    t_idx
                )


                X_tr = df_s.loc[
                    tr_idx
                ].drop(columns=TARGET_COLS)

                y_tr = df_s.loc[
                    tr_idx,
                    y_col
                ]

                X_te = df_s.loc[
                    t_idx
                ].drop(columns=TARGET_COLS)

                y_te = df_s.loc[
                    t_idx,
                    y_col
                ]


                sc_y = StandardScaler()

                y_tr_sc = sc_y.fit_transform(
                    y_tr.values.reshape(-1, 1)
                ).ravel()


                for nome, pipe in modelos.items():

                    m = clone(pipe).fit(
                        X_tr,
                        y_tr_sc
                    )

                    pr = sc_y.inverse_transform(
                        m.predict(X_te).reshape(-1, 1)
                    ).ravel()


                    metrics_rep[nome].append(
                        calc_metrics(y_te, pr)
                    )


                    for r, p in zip(y_te, pr):

                        todas_observacoes.append({
                            "Metodo": "M2",
                            "Var": y_col,
                            "Safra": s,
                            "Mod": nome,
                            "Cen": f"Rem_{n_rem}",
                            "Rep": rep + 1,
                            "Real": r,
                            "Pred": p
                        })


            for nome in modelos:

                res_m2.append({
                    "Variável": y_col,
                    "Safra": s,
                    "Modelo": nome,
                    "Rem": n_rem,
                    **pd.DataFrame(
                        metrics_rep[nome]
                    ).mean().to_dict()
                })


        # ==================================================
        # M3: Estratificado N (1-2 blocos/N nível)
        # ==================================================
        for n_rem_n in [1, 2]:

            metrics_rep = {
                m: []
                for m in modelos
            }


            for rep in range(50):

                t_idx = (
                    df_s
                    .groupby(
                        ["Variedade", "N"],
                        group_keys=False
                    )
                    .sample(
                        n=n_rem_n,
                        random_state=rep * 100
                    )
                    .index
                )


                tr_idx = df_s.index.difference(
                    t_idx
                )


                X_tr = df_s.loc[
                    tr_idx
                ].drop(columns=TARGET_COLS)

                y_tr = df_s.loc[
                    tr_idx,
                    y_col
                ]

                X_te = df_s.loc[
                    t_idx
                ].drop(columns=TARGET_COLS)

                y_te = df_s.loc[
                    t_idx,
                    y_col
                ]


                sc_y = StandardScaler()

                y_tr_sc = sc_y.fit_transform(
                    y_tr.values.reshape(-1, 1)
                ).ravel()


                for nome, pipe in modelos.items():

                    m = clone(pipe).fit(
                        X_tr,
                        y_tr_sc
                    )

                    pr = sc_y.inverse_transform(
                        m.predict(X_te).reshape(-1, 1)
                    ).ravel()


                    metrics_rep[nome].append(
                        calc_metrics(y_te, pr)
                    )


                    for r, p in zip(y_te, pr):

                        todas_observacoes.append({
                            "Metodo": "M3",
                            "Var": y_col,
                            "Safra": s,
                            "Mod": nome,
                            "Cen": f"RemN_{n_rem_n}",
                            "Rep": rep + 1,
                            "Real": r,
                            "Pred": p
                        })


            for nome in modelos:

                res_m3.append({
                    "Variável": y_col,
                    "Safra": s,
                    "Modelo": nome,
                    "RemN": n_rem_n,
                    **pd.DataFrame(
                        metrics_rep[nome]
                    ).mean().to_dict()
                })


        # ==================================================
        # M4: Cross-N
        # ==================================================
        idx_a = df_s[
            df_s["N"] == "Alto"
        ].index

        idx_b = df_s[
            df_s["N"] == "Baixo"
        ].index


        for nome, pipe in modelos.items():

            sc_a = StandardScaler()

            y_a_sc = sc_a.fit_transform(
                df_s.loc[
                    idx_a,
                    y_col
                ].values.reshape(-1, 1)
            ).ravel()


            m_a = clone(pipe).fit(
                df_s.loc[
                    idx_a
                ].drop(columns=TARGET_COLS),
                y_a_sc
            )


            pr_b = sc_a.inverse_transform(
                m_a.predict(
                    df_s.loc[
                        idx_b
                    ].drop(columns=TARGET_COLS)
                ).reshape(-1, 1)
            ).ravel()


            res_m4.append({
                "Var": y_col,
                "Safra": s,
                "Mod": nome,
                "Dir": "A->B",
                **calc_metrics(
                    df_s.loc[
                        idx_b,
                        y_col
                    ],
                    pr_b
                )
            })


            for r, p in zip(
                df_s.loc[idx_b, y_col],
                pr_b
            ):

                todas_observacoes.append({
                    "Metodo": "M4",
                    "Var": y_col,
                    "Safra": s,
                    "Mod": nome,
                    "Cen": "A->B",
                    "Rep": 1,
                    "Real": r,
                    "Pred": p
                })


# ==========================================================
# 5. SALVAMENTO DOS RESULTADOS
# ==========================================================
document.add_heading(
    "M1: Baseline Intra-Safra",
    2
)

adicionar_dataframe_ao_doc(
    document,
    pd.DataFrame(res_m1)
)


document.add_heading(
    "M2: Subfenotipagem Aleatória",
    2
)

adicionar_dataframe_ao_doc(
    document,
    pd.DataFrame(res_m2)
)


document.add_heading(
    "M3: Subfenotipagem Estratificada",
    2
)

adicionar_dataframe_ao_doc(
    document,
    pd.DataFrame(res_m3)
)


document.add_heading(
    "M4: Predição Cross-N",
    2
)

adicionar_dataframe_ao_doc(
    document,
    pd.DataFrame(res_m4)
)


document.save(
    "Relatorio_Estatisticas.docx"
)


pd.DataFrame(
    todas_observacoes
).to_csv(
    "Dados_Brutos_Predicoes.csv",
    index=False,
    sep=";",
    decimal=","
)


print(
    "\nFinalizado! "
    "'Relatorio_Estatisticas.docx' e "
    "'Dados_Brutos_Predicoes.csv' criados."
)
